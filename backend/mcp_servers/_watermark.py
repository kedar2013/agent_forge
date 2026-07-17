"""Shared "Agent Forge" diagonal watermark applied to every
generated report — PDF (document_export_server.py), Excel
(document_export_server.py), and PPTX (slide_reporting_server.py) — so a
file downloaded from chat is still traceable back to the app it came from
once it's off-platform (forwarded, printed, re-uploaded elsewhere).

Each format needs its own rendering technique (reportlab canvas rotation
for PDF, a composited PNG for Excel since openpyxl has no native diagonal
text, raw XML alpha for PPTX since python-pptx's high-level Font API has
no transparency setter) — this module is just the one place the shared
text/color/alpha constants live, so all three stay visually consistent.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from PIL import Image as PILImage, ImageDraw, ImageFont

if TYPE_CHECKING:
    from reportlab.pdfgen.canvas import Canvas

WATERMARK_TEXT = "Agent Forge"

# Brand purple (matches THEME["brand"] / document_export_server.py's PDF
# title and table-header color) kept faint enough not to compete with
# actual content.
_WATERMARK_RGB = (0x5B, 0x3F, 0xE6)
_PDF_ALPHA = 0.09
_EXCEL_ALPHA = 40  # out of 255 (~16%) -- rendered as a raster overlay, so a
# touch stronger than the PDF/PPTX vector alpha to stay legible once Excel
# composites it over the grid.
_PPTX_ALPHA_PCT = 9000  # out of 100000 (~9%), OOXML's <a:alpha val=".."/> unit

_FONT_CANDIDATES = ("arialbd.ttf", "Arial Bold.ttf", "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf")


def _load_font(size: int) -> ImageFont.FreeTypeFont:
    for name in _FONT_CANDIDATES:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def watermark_pdf_page(canvas: "Canvas", page_width: float, page_height: float) -> None:
    """Draws the diagonal watermark on the current reportlab page. Call this
    as (or from) SimpleDocTemplate.build's onFirstPage/onLaterPages callback
    — those fire once per page with the live canvas, before the page's own
    content is flushed, so the watermark ends up correctly behind it."""
    from reportlab.lib.colors import Color

    canvas.saveState()
    canvas.setFillColor(Color(_WATERMARK_RGB[0] / 255, _WATERMARK_RGB[1] / 255, _WATERMARK_RGB[2] / 255, alpha=_PDF_ALPHA))
    canvas.setFont("Helvetica-Bold", 40)
    canvas.translate(page_width / 2, page_height / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, 0, WATERMARK_TEXT)
    canvas.restoreState()


def render_excel_watermark_image(width_px: int, height_px: int) -> io.BytesIO:
    """A transparent PNG, in-memory (no temp file needed — openpyxl's Image
    accepts a file-like object) with the watermark text rotated and
    centered, sized to roughly cover a sheet's used range. Excel has no
    native diagonal-text watermark, and a floating image always draws on
    top of cell content (there's no z-order trick to put it "behind" the
    grid) — kept faint enough via alpha that data underneath stays readable.

    Returned as a BytesIO rather than the raw PIL Image: openpyxl's
    Image._data() assumes `self.ref.fp` exists (a real file handle), which
    only PIL images opened FROM a file/stream have — a freshly composited
    in-memory image has no .fp of its own. Round-tripping through a buffer
    (save PNG bytes -> reopen) gives openpyxl a `.fp` it can read back."""
    width_px = max(500, min(width_px, 2000))
    height_px = max(300, min(height_px, 1400))

    font = _load_font(48)
    probe = ImageDraw.Draw(PILImage.new("RGBA", (1, 1)))
    left, top, right, bottom = probe.textbbox((0, 0), WATERMARK_TEXT, font=font)
    text_w, text_h = right - left, bottom - top
    pad = 24
    text_layer = PILImage.new("RGBA", (text_w + pad * 2, text_h + pad * 2), (0, 0, 0, 0))
    ImageDraw.Draw(text_layer).text((pad, pad), WATERMARK_TEXT, font=font, fill=(*_WATERMARK_RGB, _EXCEL_ALPHA))
    rotated = text_layer.rotate(35, expand=True, resample=PILImage.BICUBIC)

    canvas = PILImage.new("RGBA", (width_px, height_px), (0, 0, 0, 0))
    x = (width_px - rotated.width) // 2
    y = (height_px - rotated.height) // 2
    canvas.alpha_composite(rotated, (x, y))

    buf = io.BytesIO()
    canvas.save(buf, format="PNG")
    buf.seek(0)
    return buf


def watermark_pptx_slide(slide, slide_width_emu: int, slide_height_emu: int) -> None:
    """Adds a diagonal, low-alpha textbox spanning the slide. Inserted as
    the FIRST shape on the slide (call before adding any other content) —
    shape z-order in a .pptx follows document order, so the watermark ends
    up rendered behind everything the slide's own builder adds afterward.
    Alpha needs raw XML (python-pptx's Font.color has no transparency
    setter) — `_apply_run_alpha` reaches into the run's <a:srgbClr> right
    after python-pptx creates it via the normal `.rgb =` assignment."""
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN

    box_w, box_h = Emu(int(slide_width_emu * 1.3)), Emu(Pt(70))
    left = Emu(int((slide_width_emu - box_w) / 2))
    top = Emu(int((slide_height_emu - box_h) / 2))

    tb = slide.shapes.add_textbox(left, top, box_w, box_h)
    tb.rotation = -35
    tf = tb.text_frame
    # Without this, PowerPoint auto-shrinks the box to fit the text at
    # render time (python-pptx's textbox default) — which would throw off
    # the centering computed above from the box's original dimensions.
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = WATERMARK_TEXT
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*_WATERMARK_RGB)

    rPr = run.font._rPr
    solid_fill = rPr.find(qn("a:solidFill"))
    srgb_clr = solid_fill.find(qn("a:srgbClr")) if solid_fill is not None else None
    if srgb_clr is not None:
        alpha_el = etree.SubElement(srgb_clr, qn("a:alpha"))
        alpha_el.set("val", str(_PPTX_ALPHA_PCT))

    # Moved to the front of the slide's shape tree so later shapes (heading,
    # chart, table, ...) render on top of it instead of the other way round.
    sp_tree = tb._element.getparent()
    sp_tree.remove(tb._element)
    sp_tree.insert(2, tb._element)  # after <p:nvGrpSpPr/> and <p:grpSpPr/>
